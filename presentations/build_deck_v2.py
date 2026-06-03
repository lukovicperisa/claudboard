"""
Generates the team presentation, themed to match the claudboard UI mockup
(/Users/LUP1BG/Documents/BoschProjects/bosch-workflow/Bosch workflow/).

Design tokens lifted directly from src/styles.css:
  --bg #08090a · --surface #131418 · --border #23252b
  --teal (primary), --green (success), --amber (gate), --violet (prep)
  Geist Sans / Geist Mono typography
"""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

OUT = Path(__file__).parent / "AI-SDLC-team-presentation-v2.pptx"
SHOTS = Path(__file__).parent / "screenshots"

# ---- Tokens from styles.css ---------------------------------------------
BG          = RGBColor(0x08, 0x09, 0x0A)
BG_2        = RGBColor(0x0E, 0x0F, 0x11)
SURFACE     = RGBColor(0x13, 0x14, 0x18)
SURFACE_2   = RGBColor(0x18, 0x1A, 0x1F)
SURFACE_3   = RGBColor(0x20, 0x23, 0x2A)
BORDER      = RGBColor(0x23, 0x25, 0x2B)
BORDER_HI   = RGBColor(0x2E, 0x31, 0x38)
BORDER_BR   = RGBColor(0x3B, 0x3F, 0x48)

TEXT        = RGBColor(0xED, 0xED, 0xEF)
TEXT_2      = RGBColor(0xB8, 0xB9, 0xBE)
MUTED       = RGBColor(0x79, 0x7B, 0x83)
DIM         = RGBColor(0x4D, 0x4F, 0x57)

# Accent colors — oklch converted to approximate RGB
TEAL        = RGBColor(0x4D, 0xD5, 0xDA)
TEAL_DIM    = RGBColor(0x14, 0x28, 0x2A)   # ~teal at 18% on dark
GREEN       = RGBColor(0x5B, 0xD8, 0x9A)
GREEN_DIM   = RGBColor(0x13, 0x29, 0x1F)
AMBER       = RGBColor(0xE3, 0xB0, 0x54)
AMBER_DIM   = RGBColor(0x2C, 0x21, 0x10)
RED         = RGBColor(0xE3, 0x66, 0x58)
RED_DIM     = RGBColor(0x2C, 0x14, 0x10)
VIOLET      = RGBColor(0xA7, 0x8B, 0xF0)
VIOLET_DIM  = RGBColor(0x20, 0x18, 0x2E)

# On-accent text (for solid backgrounds, like .btn.accent)
ON_TEAL     = BG_2
ON_AMBER    = RGBColor(0x1A, 0x13, 0x00)
ON_GREEN    = BG_2

FONT_SANS   = "Inter"        # Geist fallback chain → Inter is closest commonly installed
FONT_MONO   = "Menlo"        # Geist Mono fallback

# ---- Presentation -------------------------------------------------------
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]


# ---- Primitives ---------------------------------------------------------
def add_rect(slide, x, y, w, h, fill, line=None, line_w=0.5):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(line_w)
    shape.shadow.inherit = False
    return shape


def add_round(slide, x, y, w, h, fill, line=None, line_w=0.5, radius=0.06):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    shape.adjustments[0] = radius
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(line_w)
    shape.shadow.inherit = False
    return shape


def add_oval(slide, x, y, w, h, fill, line=None, line_w=0.5):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(line_w)
    shape.shadow.inherit = False
    return shape


def add_text(slide, x, y, w, h, text, *, size=14, bold=False, color=TEXT,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, font=FONT_SANS,
             line_spacing=1.25):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.04)
    tf.margin_right = Inches(0.04)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = anchor
    lines = text.split("\n") if isinstance(text, str) else text
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        run = p.add_run()
        run.text = line
        run.font.name = font
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
    return tb


def add_bullets(slide, x, y, w, h, bullets, *, size=13, color=TEXT_2,
                line_spacing=1.45, marker_color=TEAL, marker="—",
                marker_font=FONT_MONO):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.04)
    tf.margin_top = Inches(0.02)
    for i, b in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.line_spacing = line_spacing
        p.space_after = Pt(5)
        dot = p.add_run()
        dot.text = f"{marker}  "
        dot.font.name = marker_font
        dot.font.size = Pt(size)
        dot.font.bold = False
        dot.font.color.rgb = marker_color
        body = p.add_run()
        body.text = b
        body.font.name = FONT_SANS
        body.font.size = Pt(size)
        body.font.color.rgb = color
    return tb


def add_chip(slide, x, y, w, h, label, *, fill=SURFACE, fg=TEXT_2,
             border=BORDER_HI, size=10, mono=False, bold=False):
    add_round(slide, x, y, w, h, fill, line=border, line_w=0.5, radius=0.5)
    add_text(slide, x, y, w, h, label, size=size, bold=bold, color=fg,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
             font=FONT_MONO if mono else FONT_SANS)


def add_chip_dot(slide, x, y, w, h, label, *, dim_fill, dot_color,
                 fg, size=10):
    """Chip with a leading colored dot — mimics .chip with .dot."""
    add_round(slide, x, y, w, h, dim_fill, line=None, radius=0.5)
    add_oval(slide, x + Inches(0.1), y + (h - Inches(0.08)) / 2,
             Inches(0.08), Inches(0.08), dot_color)
    add_text(slide, x + Inches(0.22), y, w - Inches(0.25), h, label,
             size=size, color=fg, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.MIDDLE, font=FONT_SANS, bold=True)


def add_card(slide, x, y, w, h, *, fill=SURFACE, border=BORDER, radius=0.04):
    add_round(slide, x, y, w, h, fill, line=border, line_w=0.5, radius=radius)


def add_brand(slide, x, y, *, sub="v3.8.1", mark="cb"):
    """Sidebar-style brand mark + wordmark."""
    add_round(slide, x, y, Inches(0.32), Inches(0.32), TEXT, radius=0.18)
    add_text(slide, x, y, Inches(0.32), Inches(0.32), mark,
             size=11, bold=True, color=BG, font=FONT_MONO,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_text(slide, x + Inches(0.42), y + Inches(0.02), Inches(2.0),
             Inches(0.3), "claudboard",
             size=15, bold=True, color=TEXT, font=FONT_SANS)
    # version pill
    add_round(slide, x + Inches(1.55), y + Inches(0.06), Inches(0.55),
              Inches(0.22), BG, line=BORDER_HI, line_w=0.5, radius=0.5)
    add_text(slide, x + Inches(1.55), y + Inches(0.05), Inches(0.55),
             Inches(0.22), sub.upper(), size=8, color=MUTED,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
             font=FONT_MONO, bold=True)


def add_eyebrow(slide, x, y, w, text, *, color=MUTED):
    add_text(slide, x, y, w, Inches(0.3), text.upper(),
             size=10, bold=True, color=color, font=FONT_SANS)


def page_header(slide, eyebrow, title, *, sub=None):
    add_brand(slide, Inches(0.55), Inches(0.42))
    # right-side path crumb like topbar
    add_text(slide, Inches(8.0), Inches(0.48), Inches(5.2), Inches(0.3),
             "~/claude-repo-scan  ·  main",
             size=10, color=DIM, font=FONT_MONO, align=PP_ALIGN.RIGHT)
    # hairline divider below brand row
    add_rect(slide, Inches(0.55), Inches(0.95), Inches(12.2), Emu(6000), BORDER)

    add_eyebrow(slide, Inches(0.55), Inches(1.2), Inches(8), eyebrow,
                color=TEAL)
    add_text(slide, Inches(0.55), Inches(1.5), Inches(12.2), Inches(0.7),
             title, size=28, bold=True, color=TEXT,
             line_spacing=1.1)
    if sub:
        add_text(slide, Inches(0.55), Inches(2.05), Inches(12.2),
                 Inches(0.4), sub, size=13, color=MUTED)


def footer(slide, page_text):
    add_rect(slide, Inches(0.55), Inches(7.05), Inches(12.2), Emu(6000),
             BORDER)
    add_oval(slide, Inches(0.55), Inches(7.15), Inches(0.1), Inches(0.1),
             TEAL)
    add_text(slide, Inches(0.72), Inches(7.13), Inches(8), Inches(0.3),
             "claude-repo-scan  ·  bosch-sdlc  ·  team review",
             size=9, color=MUTED, font=FONT_SANS)
    add_text(slide, Inches(11.0), Inches(7.13), Inches(2.0), Inches(0.3),
             page_text, size=9, color=DIM, font=FONT_MONO,
             align=PP_ALIGN.RIGHT)


def blank_slide():
    s = prs.slides.add_slide(BLANK)
    add_rect(s, 0, 0, SW, SH, BG)
    return s


# =========================================================================
# Slide 1 — Title (hero)
# =========================================================================
s = blank_slide()
# faint dot grid background
for r in range(11):
    for c in range(26):
        add_oval(s, Inches(0.3 + c * 0.5),
                 Inches(0.4 + r * 0.4),
                 Inches(0.03), Inches(0.03), SURFACE_3)

# brand top-left
add_brand(s, Inches(0.6), Inches(0.6))

# right-side meta
add_text(s, Inches(8.0), Inches(0.65), Inches(5.0), Inches(0.3),
         "TEAM REVIEW  ·  50-DAY BUILD",
         size=10, bold=True, color=MUTED, font=FONT_SANS,
         align=PP_ALIGN.RIGHT)

# Title block
add_text(s, Inches(0.6), Inches(2.4), Inches(12), Inches(1.1),
         "AUTONOMOUS",
         size=12, bold=True, color=TEAL, font=FONT_MONO)

add_text(s, Inches(0.6), Inches(2.75), Inches(12), Inches(1.5),
         "Feature delivery,", size=58, bold=True, color=TEXT,
         line_spacing=1.0)
add_text(s, Inches(0.6), Inches(3.75), Inches(12), Inches(1.5),
         "with one human gate.", size=58, bold=True, color=TEXT_2,
         line_spacing=1.0)

# subtitle
add_text(s, Inches(0.6), Inches(5.05), Inches(12), Inches(0.5),
         "Two components. Same workflow. Any Bosch repo.",
         size=18, color=MUTED, font=FONT_SANS)

# chips strip (like .chip)
chip_y = Inches(5.95)
chips = [
    ("vertical",   "claudboard plugin",       VIOLET, VIOLET_DIM),
    ("horizontal", "feature-workflow skill",  TEAL,   TEAL_DIM),
    ("surface",    "Claude CLI  ·  SDLC UI",  AMBER,  AMBER_DIM),
]
cx = Inches(0.6)
for kind, text, color, dim in chips:
    chip_w = Inches(0.18 + 0.085 * len(text) + 0.85)
    add_round(s, cx, chip_y, chip_w, Inches(0.4), dim, radius=0.5)
    add_text(s, cx + Inches(0.18), chip_y, Inches(1.0), Inches(0.4),
             kind, size=10, bold=True, color=color, font=FONT_MONO,
             anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, cx + Inches(1.05), chip_y, chip_w - Inches(1.15),
             Inches(0.4), text, size=12, color=TEXT,
             font=FONT_SANS, anchor=MSO_ANCHOR.MIDDLE)
    cx += chip_w + Inches(0.15)

# footer meta
add_rect(s, Inches(0.6), Inches(6.85), Inches(12.2), Emu(6000), BORDER)
add_text(s, Inches(0.6), Inches(6.95), Inches(12), Inches(0.4),
         "AI SDLC INITIATIVE", size=10, bold=True, color=DIM,
         font=FONT_MONO)
add_text(s, Inches(8.0), Inches(6.95), Inches(5.0), Inches(0.4),
         "claude-repo-scan  +  bosch-sdlc",
         size=10, color=DIM, font=FONT_MONO, align=PP_ALIGN.RIGHT)


# =========================================================================
# Slide 2 — The problem
# =========================================================================
s = blank_slide()
page_header(s, "Why we built this",
            "A typical feature today.",
            sub="Three high-touch phases. Most of it is repetitive context-shuffling.")

# Three "prereq"-style cards (like .prereq pattern)
y = Inches(2.8)
card_w = Inches(3.95)
card_h = Inches(2.3)
gap = Inches(0.2)
x0 = Inches(0.55)

steps = [
    ("01", "Refine",
     "Ticket grooming · clarification\nback-and-forth with stakeholders.",
     VIOLET, VIOLET_DIM),
    ("02", "Specify",
     "BDD / acceptance criteria · design\narchitecture · plan decomposition.",
     TEAL, TEAL_DIM),
    ("03", "Implement",
     "Branch · code · test · review\ncommit · open PR · finalize ticket.",
     AMBER, AMBER_DIM),
]
for i, (num, title, body, color, dim) in enumerate(steps):
    x = x0 + (card_w + gap) * i
    add_card(s, x, y, card_w, card_h)
    # numbered icon block (like .pico)
    add_round(s, x + Inches(0.3), y + Inches(0.35), Inches(0.5),
              Inches(0.5), dim, radius=0.18)
    add_text(s, x + Inches(0.3), y + Inches(0.35), Inches(0.5),
             Inches(0.5), num, size=12, bold=True, color=color,
             font=FONT_MONO, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    # title
    add_text(s, x + Inches(0.95), y + Inches(0.35), Inches(2.5),
             Inches(0.4), title, size=18, bold=True, color=TEXT)
    add_text(s, x + Inches(0.95), y + Inches(0.7), Inches(2.5),
             Inches(0.3), "cycle", size=10, color=MUTED, font=FONT_MONO)
    # divider
    add_rect(s, x + Inches(0.3), y + Inches(1.15), card_w - Inches(0.6),
             Emu(6000), BORDER)
    add_text(s, x + Inches(0.3), y + Inches(1.3), card_w - Inches(0.6),
             Inches(1.0), body, size=12, color=TEXT_2,
             line_spacing=1.45)

# Insight band (matches .run-banner pattern with gradient feel)
band_y = Inches(5.5)
add_card(s, Inches(0.55), band_y, Inches(12.2), Inches(1.4))
# Left accent strip (mimics gradient banner)
add_rect(s, Inches(0.55), band_y, Inches(0.04), Inches(1.4), TEAL)
# Icon block
add_round(s, Inches(0.85), band_y + Inches(0.35), Inches(0.7),
          Inches(0.7), TEAL_DIM, radius=0.15)
add_text(s, Inches(0.85), band_y + Inches(0.35), Inches(0.7),
         Inches(0.7), "→", size=24, bold=True, color=TEAL,
         font=FONT_MONO, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
# Text
add_eyebrow(s, Inches(1.75), band_y + Inches(0.25), Inches(8), "OUR BET",
            color=TEAL)
add_text(s, Inches(1.75), band_y + Inches(0.55), Inches(10.5),
         Inches(0.75),
         "Most of this can run autonomously — if Claude knows the project\n"
         "and follows a deterministic, instrumented workflow.",
         size=16, bold=True, color=TEXT, line_spacing=1.3)

footer(s, "02 / 13")


# =========================================================================
# Slide 3 — Overview: two components
# =========================================================================
s = blank_slide()
page_header(s, "Overview",
            "Two components. One goal.",
            sub="The plugin makes Claude project-aware. The workflow uses that awareness to ship.")

# Left card — Vertical
lc_x = Inches(0.55); lc_y = Inches(2.8)
lc_w = Inches(6.0);  lc_h = Inches(4.1)
add_card(s, lc_x, lc_y, lc_w, lc_h)

# Card head (.card-head pattern)
add_rect(s, lc_x + Inches(0.3), lc_y + Inches(0.85),
         lc_w - Inches(0.6), Emu(6000), BORDER)
# kind chip
add_round(s, lc_x + Inches(0.3), lc_y + Inches(0.3),
          Inches(1.2), Inches(0.35), VIOLET_DIM, radius=0.5)
add_text(s, lc_x + Inches(0.3), lc_y + Inches(0.3),
         Inches(1.2), Inches(0.35), "VERTICAL",
         size=10, bold=True, color=VIOLET, font=FONT_MONO,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
add_text(s, lc_x + Inches(1.6), lc_y + Inches(0.3),
         Inches(4), Inches(0.35), "project context",
         size=12, color=MUTED, anchor=MSO_ANCHOR.MIDDLE, font=FONT_MONO)
# title
add_text(s, lc_x + Inches(0.3), lc_y + Inches(1.05),
         lc_w - Inches(0.6), Inches(0.55), "claudboard plugin",
         size=22, bold=True, color=TEXT)
add_text(s, lc_x + Inches(0.3), lc_y + Inches(1.65),
         lc_w - Inches(0.6), Inches(0.4),
         "Teaches Claude any repo in minutes.",
         size=13, color=TEXT_2)
add_bullets(s, lc_x + Inches(0.3), lc_y + Inches(2.25),
            lc_w - Inches(0.6), lc_h - Inches(2.4), [
    "/analyse    scan stack, patterns, conventions",
    "/generate   CLAUDE.md, rules, scoped skills",
    "/refresh    keep artifacts in sync with code",
    "/techdebt   ticket-ready refactor backlog",
], size=13, marker_color=VIOLET, marker_font=FONT_MONO)

# Right card — Horizontal
rc_x = Inches(6.75); rc_y = Inches(2.8)
rc_w = Inches(6.0);  rc_h = Inches(4.1)
add_card(s, rc_x, rc_y, rc_w, rc_h)
add_rect(s, rc_x + Inches(0.3), rc_y + Inches(0.85),
         rc_w - Inches(0.6), Emu(6000), BORDER)
add_round(s, rc_x + Inches(0.3), rc_y + Inches(0.3),
          Inches(1.4), Inches(0.35), TEAL_DIM, radius=0.5)
add_text(s, rc_x + Inches(0.3), rc_y + Inches(0.3),
         Inches(1.4), Inches(0.35), "HORIZONTAL",
         size=10, bold=True, color=TEAL, font=FONT_MONO,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
add_text(s, rc_x + Inches(1.8), rc_y + Inches(0.3),
         Inches(4), Inches(0.35), "feature delivery",
         size=12, color=MUTED, anchor=MSO_ANCHOR.MIDDLE, font=FONT_MONO)
add_text(s, rc_x + Inches(0.3), rc_y + Inches(1.05),
         rc_w - Inches(0.6), Inches(0.55), "feature-workflow skill",
         size=22, bold=True, color=TEXT)
add_text(s, rc_x + Inches(0.3), rc_y + Inches(1.65),
         rc_w - Inches(0.6), Inches(0.4),
         "Runs the full feature lifecycle, autonomously.",
         size=13, color=TEXT_2)
add_bullets(s, rc_x + Inches(0.3), rc_y + Inches(2.25),
            rc_w - Inches(0.6), rc_h - Inches(2.4), [
    "7 phases — ticket → PR → finalize",
    "10 specialized sub-agents",
    "1 human gate (spec + plan)",
    "Runs in Claude CLI or in the Bosch SDLC UI",
], size=13, marker_color=TEAL, marker_font=FONT_MONO)

footer(s, "03 / 13")


# =========================================================================
# Slide 4 — Plugin deep-dive (matches .prereq layout)
# =========================================================================
s = blank_slide()
page_header(s, "Component 1  ·  Vertical",
            "claudboard — onboarding any repo.",
            sub="Four sub-commands. Each one is a full-scope skill, not a stub.")

cards = [
    ("/analyse",  "Read-only deep scan",
     "Detects stack, frameworks,\narchitecture, conventions,\nCI/CD, anti-patterns.",
     TEAL, TEAL_DIM),
    ("/generate", "Author .claude artifacts",
     "CLAUDE.md, rules with paths,\nfull-scope skills with refs\nand scripts.",
     GREEN, GREEN_DIM),
    ("/refresh",  "Delta updates",
     "Re-scans, diffs against\nexisting artifacts, updates\nonly what changed.",
     AMBER, AMBER_DIM),
    ("/techdebt", "Refactor backlog",
     "Module-grouped, ticket-ready\nreport with severity, effort,\nand fix suggestions.",
     VIOLET, VIOLET_DIM),
]

card_w = Inches(2.95)
card_h = Inches(3.5)
gap = Inches(0.15)
x0 = Inches(0.55)
y = Inches(2.75)

for i, (cmd, head, body, color, dim) in enumerate(cards):
    x = x0 + (card_w + gap) * i
    add_card(s, x, y, card_w, card_h)
    # icon block at top
    add_round(s, x + Inches(0.3), y + Inches(0.3), Inches(0.55),
              Inches(0.55), dim, radius=0.15)
    add_text(s, x + Inches(0.3), y + Inches(0.3), Inches(0.55),
             Inches(0.55), "▶", size=14, bold=True, color=color,
             font=FONT_MONO, align=PP_ALIGN.CENTER,
             anchor=MSO_ANCHOR.MIDDLE)
    # command (mono)
    add_text(s, x + Inches(1.0), y + Inches(0.35), card_w - Inches(1.2),
             Inches(0.4), cmd, size=14, bold=True, color=TEXT,
             font=FONT_MONO)
    add_text(s, x + Inches(1.0), y + Inches(0.7), card_w - Inches(1.2),
             Inches(0.3), "command", size=9, color=MUTED, font=FONT_MONO)
    # divider
    add_rect(s, x + Inches(0.3), y + Inches(1.15),
             card_w - Inches(0.6), Emu(6000), BORDER)
    # heading + body
    add_text(s, x + Inches(0.3), y + Inches(1.3),
             card_w - Inches(0.6), Inches(0.4), head,
             size=13, bold=True, color=TEXT)
    add_text(s, x + Inches(0.3), y + Inches(1.75),
             card_w - Inches(0.6), card_h - Inches(2.1),
             body, size=12, color=TEXT_2, line_spacing=1.5)
    # bottom status chip
    add_chip(s, x + Inches(0.3), y + card_h - Inches(0.55),
             Inches(1.1), Inches(0.3),
             "full-scope", fill=dim, fg=color, border=None, size=9,
             bold=True)

# Workspace note
nb_y = Inches(6.45)
add_card(s, Inches(0.55), nb_y, Inches(12.2), Inches(0.5))
add_rect(s, Inches(0.55), nb_y, Inches(0.04), Inches(0.5), VIOLET)
add_text(s, Inches(0.85), nb_y, Inches(12), Inches(0.5),
         "Workspace mode  ·  one meta-repo holds shared .claude/ for N service repos.",
         size=12, color=TEXT_2, anchor=MSO_ANCHOR.MIDDLE)

footer(s, "04 / 13")


# =========================================================================
# Slide 5 — Workflow at a glance
# =========================================================================
s = blank_slide()
page_header(s, "Component 2  ·  Horizontal",
            "feature-workflow — 7 phases, 1 gate.",
            sub="Phase 1 produces a spec and a plan. You approve them once. Everything else is autonomous.")

phases = [
    ("1", "Ticket\nClarify\nSpec · Plan", AMBER, AMBER_DIM, "gate"),
    ("2", "Branch", GREEN, GREEN_DIM, "auto"),
    ("3", "Develop\n& Test", GREEN, GREEN_DIM, "auto"),
    ("4", "Commit", GREEN, GREEN_DIM, "auto"),
    ("5", "Review", GREEN, GREEN_DIM, "auto"),
    ("6", "Pull\nRequest", GREEN, GREEN_DIM, "auto"),
    ("7", "Finalize", GREEN, GREEN_DIM, "auto"),
]

# Pipeline area
pipe_y = Inches(2.9)
diam = Inches(0.95)
n = len(phases)
total_w = Inches(12.2)
spacing = (total_w - diam) / (n - 1)
x_start = Inches(0.55) + (total_w - (spacing * (n - 1) + diam)) / 2

# Connecting line (hairline)
line_y = pipe_y + diam / 2
add_rect(s, x_start + diam / 2, line_y,
         spacing * (n - 1), Emu(6000), BORDER_HI)

for i, (num, title, color, dim, kind) in enumerate(phases):
    cx = x_start + spacing * i
    # outer ring
    add_oval(s, cx - Inches(0.06), pipe_y - Inches(0.06),
             diam + Inches(0.12), diam + Inches(0.12), dim)
    # main circle
    add_oval(s, cx, pipe_y, diam, diam, BG)
    add_oval(s, cx + Inches(0.02), pipe_y + Inches(0.02),
             diam - Inches(0.04), diam - Inches(0.04), dim)
    add_text(s, cx, pipe_y + Inches(0.18), diam, Inches(0.6),
             num, size=28, bold=True, color=color,
             font=FONT_MONO, align=PP_ALIGN.CENTER)
    # title under
    add_text(s, cx - Inches(0.4), pipe_y + diam + Inches(0.15),
             diam + Inches(0.8), Inches(0.9),
             title, size=11, bold=True, color=TEXT,
             align=PP_ALIGN.CENTER, line_spacing=1.2)
    # kind chip (mini)
    add_round(s, cx - Inches(0.25), pipe_y + diam + Inches(1.0),
              diam + Inches(0.5), Inches(0.28), dim, radius=0.5)
    add_text(s, cx - Inches(0.25), pipe_y + diam + Inches(1.0),
             diam + Inches(0.5), Inches(0.28), kind,
             size=8, bold=True, color=color, font=FONT_MONO,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

# Human gate banner above phase 1
gate_x = x_start
add_text(s, gate_x - Inches(0.5), pipe_y - Inches(0.6),
         diam + Inches(1), Inches(0.3),
         "▲  HUMAN GATE", size=10, bold=True, color=AMBER,
         font=FONT_MONO, align=PP_ALIGN.CENTER)

# Sub-agents row (chips) — wraps to 2 rows to stay within slide width
agt_y = Inches(5.3)
add_text(s, Inches(0.55), agt_y, Inches(12), Inches(0.3),
         "10 SPECIALIZED SUB-AGENTS", size=10, bold=True, color=MUTED,
         font=FONT_MONO)
add_text(s, Inches(0.55), agt_y + Inches(0.3), Inches(12), Inches(0.4),
         "Each one gets only the tools it needs and a self-contained prompt.",
         size=12, color=TEXT_2)

# Agent chips with dots (style = chip with .dot)
agents = [
    ("sdd-expert",      TEAL, TEAL_DIM),
    ("architect",       TEAL, TEAL_DIM),
    ("jira-agent",      VIOLET, VIOLET_DIM),
    ("tr-agent",        VIOLET, VIOLET_DIM),
    ("pr-agent-ado",    VIOLET, VIOLET_DIM),
    ("pr-agent-github", VIOLET, VIOLET_DIM),
    ("git-agent",       GREEN, GREEN_DIM),
    ("implementation",  GREEN, GREEN_DIM),
    ("spec-reviewer",   AMBER, AMBER_DIM),
    ("design-reviewer", AMBER, AMBER_DIM),
]
left_margin = Inches(0.55)
right_edge  = Inches(13.333) - Inches(0.55)   # 12.78" usable
chip_y0     = agt_y + Inches(0.9)
chip_h      = Inches(0.34)
chip_gap    = Inches(0.08)
row_gap     = Inches(0.1)
# Conservative width estimate — 10pt Inter bold runs ~0.10"/char rendered;
# base accounts for dot + padding (~0.55" total chrome per chip)
char_w      = 0.10
base_w      = 0.55

cx_cursor = left_margin
cy_cursor = chip_y0
for label, color, dim in agents:
    chip_w = Inches(base_w + char_w * len(label))
    if cx_cursor + chip_w > right_edge:
        cx_cursor = left_margin
        cy_cursor = cy_cursor + chip_h + row_gap
    add_chip_dot(s, cx_cursor, cy_cursor, chip_w, chip_h, label,
                 dim_fill=dim, dot_color=color, fg=color, size=10)
    cx_cursor += chip_w + chip_gap

footer(s, "05 / 13")


# =========================================================================
# Phase-detail slide builder (matches .pane + .phase patterns)
# =========================================================================
def phase_detail(num_label, eyebrow, title, sub,
                 lhs_title, lhs_bullets,
                 rhs_title, rhs_bullets,
                 callout=None, accent=TEAL, dim_fill=TEAL_DIM, page=""):
    s = blank_slide()
    page_header(s, eyebrow, title, sub=sub)

    # phase status badge top-right (like .chip)
    badge_x = Inches(10.7); badge_y = Inches(1.2)
    add_round(s, badge_x, badge_y, Inches(2.05), Inches(0.35),
              dim_fill, radius=0.5)
    add_oval(s, badge_x + Inches(0.15), badge_y + Inches(0.13),
             Inches(0.1), Inches(0.1), accent)
    add_text(s, badge_x + Inches(0.3), badge_y, Inches(1.7),
             Inches(0.35), num_label.lower(),
             size=10, bold=True, color=accent, font=FONT_MONO,
             anchor=MSO_ANCHOR.MIDDLE)

    # Two column cards
    col_w = Inches(6.0); col_h = Inches(4.05)
    y = Inches(2.8)

    # LHS card
    add_card(s, Inches(0.55), y, col_w, col_h)
    # head (.card-head)
    add_rect(s, Inches(0.55) + Inches(0.3), y + Inches(0.85),
             col_w - Inches(0.6), Emu(6000), BORDER)
    # numbered icon
    add_round(s, Inches(0.55) + Inches(0.3), y + Inches(0.3),
              Inches(0.4), Inches(0.4), dim_fill, radius=0.18)
    add_text(s, Inches(0.55) + Inches(0.3), y + Inches(0.3),
             Inches(0.4), Inches(0.4), "→", size=14, bold=True,
             color=accent, font=FONT_MONO, align=PP_ALIGN.CENTER,
             anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, Inches(0.55) + Inches(0.85), y + Inches(0.32),
             col_w - Inches(1.2), Inches(0.4), lhs_title,
             size=14, bold=True, color=TEXT)
    add_bullets(s, Inches(0.55) + Inches(0.3), y + Inches(1.05),
                col_w - Inches(0.6), col_h - Inches(1.15),
                lhs_bullets, size=12.5, marker_color=accent,
                color=TEXT_2)

    # RHS card
    rx = Inches(6.75)
    add_card(s, rx, y, col_w, col_h)
    add_rect(s, rx + Inches(0.3), y + Inches(0.85),
             col_w - Inches(0.6), Emu(6000), BORDER)
    add_round(s, rx + Inches(0.3), y + Inches(0.3),
              Inches(0.4), Inches(0.4), SURFACE_2, radius=0.18)
    add_text(s, rx + Inches(0.3), y + Inches(0.3),
             Inches(0.4), Inches(0.4), "i", size=14, bold=True,
             color=MUTED, font=FONT_MONO, align=PP_ALIGN.CENTER,
             anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, rx + Inches(0.85), y + Inches(0.32),
             col_w - Inches(1.2), Inches(0.4), rhs_title,
             size=14, bold=True, color=TEXT)
    add_bullets(s, rx + Inches(0.3), y + Inches(1.05),
                col_w - Inches(0.6), col_h - Inches(1.15),
                rhs_bullets, size=12.5, marker_color=MUTED,
                color=TEXT_2)

    if callout:
        cy = Inches(6.95)
        add_round(s, Inches(0.55), cy, Inches(12.2), Inches(0.35),
                  AMBER_DIM, radius=0.5)
        add_oval(s, Inches(0.85), cy + Inches(0.13),
                 Inches(0.1), Inches(0.1), AMBER)
        add_text(s, Inches(0.55), cy, Inches(12.2), Inches(0.35),
                 callout, size=10, bold=True, color=AMBER,
                 font=FONT_MONO, align=PP_ALIGN.CENTER,
                 anchor=MSO_ANCHOR.MIDDLE)

    footer(s, page)


# =========================================================================
# Slide 6 — Phase 1
# =========================================================================
phase_detail(
    "phase 1  ·  gate",
    "Workflow  ·  ticket, clarify, spec, plan",
    "Refine the idea — once, deeply.",
    "Ticket created or fetched. The user picks a clarification level. "
    "A BDD spec and an execution plan are produced and approved together.",
    "Sub-phases  1-pre  →  1d",
    [
        "1-pre · jira-agent / tr-agent — create or fetch ticket",
        "1-syn · Stated synthesis — print what we heard",
        "1a · Clarify scope — autopilot / balanced / guided / manual",
        "1b · sdd-expert-agent — Gherkin actor·action·outcome spec",
        "1c · architect-agent — execution plan with checkpoints",
        "1d · HUMAN GATE — review spec + plan, approve or return",
    ],
    "Why this design",
    [
        "One gate, not five — context isn't lost between phases",
        "Spec + plan committed with the feature, not throwaway",
        "Clarification autonomy lever — skip or get full rubric",
        "On rejection, sub-agents re-spin with user's change notes",
        "Workflow start timestamp recorded for time + cost analytics",
    ],
    callout="▲  THIS IS THE ONLY HUMAN GATE IN THE ENTIRE WORKFLOW",
    accent=AMBER, dim_fill=AMBER_DIM,
    page="06 / 13",
)

# =========================================================================
# Slide 7 — Phase 2 + 3
# =========================================================================
phase_detail(
    "phase 2 · 3  ·  auto",
    "Workflow  ·  branch and develop",
    "Branch out, then implement checkpoint by checkpoint.",
    "Once the plan is approved, the workflow opens a branch and walks the "
    "checkpoint list. Baseline verification runs in parallel with checkpoint 1.",
    "Phase 2  ·  Branch",
    [
        "git-agent creates a clean branch from main / master",
        "Naming follows project convention — feature/<ticket>/<slug>",
        "Main branch auto-detected — no per-repo config",
    ],
    "Phase 3  ·  Develop & Test",
    [
        "implementation-agent runs build → test → lint → live loops",
        "Baseline + checkpoint 1 launched IN PARALLEL (background)",
        "If baseline fails — workflow stops immediately",
        "Recoverable failures stay in the loop, ticket untouched",
        "Workspace mode: one branch name across N affected repos",
    ],
    accent=TEAL, dim_fill=TEAL_DIM,
    page="07 / 13",
)

# =========================================================================
# Slide 8 — Phase 4 + 5
# =========================================================================
phase_detail(
    "phase 4 · 5  ·  auto",
    "Workflow  ·  commit and review",
    "Clean history, then two independent reviewers.",
    "git-agent squashes checkpoint commits into one meaningful commit. "
    "Two reviewer agents check spec coverage and code quality independently.",
    "Phase 4  ·  Commit",
    [
        "git-agent  stage → squash → commit with structured message",
        "Commit message ties back to ticket and approved spec",
        "Pre-commit hooks honoured — never bypassed",
    ],
    "Phase 5  ·  Review",
    [
        "spec-reviewer — does the code satisfy every BDD scenario?",
        "design-reviewer — code quality vs repo rules + conventions",
        "On findings — implementation-agent applies fixes",
        "git-agent amends — reviewer re-runs",
        "Loop until both reviewers pass — no human in the loop",
    ],
    accent=VIOLET, dim_fill=VIOLET_DIM,
    page="08 / 13",
)

# =========================================================================
# Slide 9 — Phase 6 + 7
# =========================================================================
phase_detail(
    "phase 6 · 7  ·  auto",
    "Workflow  ·  PR and finalize",
    "Open the PR, log the work, transition the ticket.",
    "Code is pushed and a pull request is opened. The ticket is updated "
    "with the worklog, the AI cost analysis, and a success transition.",
    "Phase 6  ·  Pull Request",
    [
        "git-agent validates PR readiness, syncs and pushes",
        "pr-agent-ado / pr-agent-github creates the PR",
        "GitHub — 'Closes #N' linking + Actions run verification",
        "Workspace mode — PRs in parallel, merge order printed",
    ],
    "Phase 7  ·  Finalize",
    [
        "jira / tr-agent logs refinement + implementation time",
        "Comment with cost analysis — tokens and USD per phase",
        "Ticket transitions to its success state",
        "On failure — optional failure_transition fires automatically",
        "Full transcript persisted as JSONL — forensic record",
    ],
    accent=GREEN, dim_fill=GREEN_DIM,
    page="09 / 13",
)


# =========================================================================
# Slide 10 — CLI surface
# =========================================================================
s = blank_slide()
page_header(s, "How you run it  ·  option A",
            "Claude Code CLI — the developer path.",
            sub="Inside any onboarded repo, open Claude Code and start the workflow.")

# Terminal-style card
tx = Inches(0.55); ty = Inches(2.85); tw = Inches(8.0); th = Inches(4.1)
add_card(s, tx, ty, tw, th, fill=BG_2)
# titlebar
add_rect(s, tx, ty, tw, Inches(0.32), SURFACE_2)
# round corners hack: re-add rounded top
add_round(s, tx, ty, tw, Inches(0.32), SURFACE_2, radius=0.06)
add_rect(s, tx, ty + Inches(0.16), tw, Inches(0.16), SURFACE_2)
# traffic lights
for i, c in enumerate([RGBColor(0xFF, 0x5F, 0x57),
                       RGBColor(0xFE, 0xBC, 0x2E),
                       RGBColor(0x28, 0xC8, 0x40)]):
    add_oval(s, tx + Inches(0.14 + 0.26 * i),
             ty + Inches(0.09), Inches(0.14), Inches(0.14), c)
add_text(s, tx, ty + Inches(0.04), tw, Inches(0.24),
         "claude  ·  ~/repos/platform",
         size=10, color=MUTED, font=FONT_MONO,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
# bottom border under chrome
add_rect(s, tx, ty + Inches(0.32), tw, Emu(6000), BORDER)

# Stream content (mimics .stream)
term_lines = [
    ("$ claude",                                                              TEXT),
    ("> /start-feature Add export-as-CSV to user list",                        TEAL),
    ("",                                                                       TEXT),
    ("clarification autonomy: balanced — accept [enter] or override?",         AMBER),
    ("> [enter]",                                                              TEAL),
    ("",                                                                       TEXT),
    ("phase 1-pre  ✓  ticket PLAT-42117 created",                              GREEN),
    ("phase 1-syn  ✓  synthesis printed",                                      GREEN),
    ("phase 1a     ✓  4 dimensions clarified",                                 GREEN),
    ("phase 1b     ✓  business-behavior-spec.md",                              GREEN),
    ("phase 1c     ✓  execution-plan.md  (5 checkpoints)",                     GREEN),
    ("phase 1d     ▲  awaiting your approval (spec + plan)",                   AMBER),
]
tb = s.shapes.add_textbox(tx + Inches(0.35), ty + Inches(0.5),
                          tw - Inches(0.6), th - Inches(0.6))
tf = tb.text_frame; tf.word_wrap = True
for i, (line, color) in enumerate(term_lines):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    p.space_after = Pt(3)
    r = p.add_run()
    r.text = line if line else " "
    r.font.name = FONT_MONO
    r.font.size = Pt(12)
    r.font.color.rgb = color

# Right side — value card
rx = Inches(8.85); ry = Inches(2.85); rw = Inches(3.9); rh = Inches(4.1)
add_card(s, rx, ry, rw, rh)
# head
add_rect(s, rx + Inches(0.3), ry + Inches(0.85),
         rw - Inches(0.6), Emu(6000), BORDER)
add_round(s, rx + Inches(0.3), ry + Inches(0.3),
          Inches(0.4), Inches(0.4), TEAL_DIM, radius=0.18)
add_text(s, rx + Inches(0.3), ry + Inches(0.3),
         Inches(0.4), Inches(0.4), "❯", size=14, bold=True,
         color=TEAL, font=FONT_MONO, align=PP_ALIGN.CENTER,
         anchor=MSO_ANCHOR.MIDDLE)
add_text(s, rx + Inches(0.85), ry + Inches(0.32),
         rw - Inches(1.0), Inches(0.4),
         "for developers", size=14, bold=True, color=TEXT)
add_text(s, rx + Inches(0.3), ry + Inches(1.05),
         rw - Inches(0.6), Inches(0.4),
         "Where it shines",
         size=10, bold=True, color=MUTED, font=FONT_MONO)
add_bullets(s, rx + Inches(0.3), ry + Inches(1.5),
            rw - Inches(0.6), rh - Inches(1.7), [
    "Zero setup beyond Claude Code",
    "Full terminal feedback, fast iteration",
    "Best for the person who wrote the ticket",
    "Same workflow, same skill files",
    "Great for spikes and quick features",
], size=12.5, marker_color=TEAL)

footer(s, "10 / 13")


# =========================================================================
# Slide 11 — UI surface (real screenshot, same layout as v1 — window + card)
# =========================================================================
s = blank_slide()
page_header(s, "How you run it  ·  option B",
            "Bosch SDLC UI — the dashboard path.",
            sub="Same workflow, driven by a local web app. One command — npx bosch-sdlc.")

# App window with the real Active Run screenshot
bx = Inches(0.55); by = Inches(2.85); bw = Inches(8.0); bh = Inches(4.1)
add_card(s, bx, by, bw, bh)
# titlebar (mac-style)
add_round(s, bx, by, bw, Inches(0.32), SURFACE_2, radius=0.06)
add_rect(s, bx, by + Inches(0.16), bw, Inches(0.16), SURFACE_2)
for i, c in enumerate([RGBColor(0xFF, 0x5F, 0x57),
                       RGBColor(0xFE, 0xBC, 0x2E),
                       RGBColor(0x28, 0xC8, 0x40)]):
    add_oval(s, bx + Inches(0.14 + 0.26 * i),
             by + Inches(0.09), Inches(0.14), Inches(0.14), c)
add_text(s, bx, by + Inches(0.04), bw, Inches(0.24),
         "claudboard  ·  localhost:3001",
         size=10, color=MUTED, font=FONT_MONO,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
add_rect(s, bx, by + Inches(0.32), bw, Emu(6000), BORDER)

# Real screenshot — Active Run (sits inside the window chrome)
img_x = bx + Inches(0.03)
img_y = by + Inches(0.32)
img_w = bw - Inches(0.06)
img_h = bh - Inches(0.32) - Inches(0.03)
s.shapes.add_picture(str(SHOTS / "run.png"), img_x, img_y,
                     width=img_w, height=img_h)

# Right side — value card (identical to v1)
rx = Inches(8.85); ry = Inches(2.85); rw = Inches(3.9); rh = Inches(4.1)
add_card(s, rx, ry, rw, rh)
add_rect(s, rx + Inches(0.3), ry + Inches(0.85),
         rw - Inches(0.6), Emu(6000), BORDER)
add_round(s, rx + Inches(0.3), ry + Inches(0.3),
          Inches(0.4), Inches(0.4), AMBER_DIM, radius=0.18)
add_text(s, rx + Inches(0.3), ry + Inches(0.3),
         Inches(0.4), Inches(0.4), "◇", size=14, bold=True,
         color=AMBER, font=FONT_MONO, align=PP_ALIGN.CENTER,
         anchor=MSO_ANCHOR.MIDDLE)
add_text(s, rx + Inches(0.85), ry + Inches(0.32),
         rw - Inches(1.0), Inches(0.4),
         "for teams", size=14, bold=True, color=TEXT)
add_text(s, rx + Inches(0.3), ry + Inches(1.05),
         rw - Inches(0.6), Inches(0.4),
         "Where it shines",
         size=10, bold=True, color=MUTED, font=FONT_MONO)
add_bullets(s, rx + Inches(0.3), ry + Inches(1.5),
            rw - Inches(0.6), rh - Inches(1.7), [
    "Non-engineers can monitor + gate runs",
    "Live phase / checkpoint / agent events",
    "Pipeline, Stream, Telemetry panes",
    "Reuses your Claude Code MCP config",
    "Per-run JSONL transcripts on disk",
], size=12.5, marker_color=AMBER)

footer(s, "11 / 13")


# =========================================================================
# Slide 12 — Reusability
# =========================================================================
s = blank_slide()
page_header(s, "Built for reuse",
            "One toolchain — every Bosch project.",
            sub="The plugin generates the right artifacts for whichever topology you have.")

topos = [
    ("monolith",  "Single repo",
     "Plugin generates .claude/ in the repo.\nWorkflow lives at\n.claude/skills/feature-workflow/.\n\n1 ticket → 1 branch → 1 PR.",
     TEAL, TEAL_DIM),
    ("monorepo",  "One repo, many modules",
     "Same install as a single repo.\nNo scope picker — agents pick\naffected modules from the prompt.\n\n1 ticket → 1 branch → 1 PR.",
     VIOLET, VIOLET_DIM),
    ("workspace", "N independent repos",
     "Workspace meta-repo holds the\nshared .claude/.\n\n1 ticket → N branches → N PRs.\nMerge order printed at the end.",
     AMBER, AMBER_DIM),
]

y = Inches(2.85)
card_w = Inches(3.95)
card_h = Inches(3.6)
gap = Inches(0.2)
x0 = Inches(0.55)

for i, (eyebrow, title, body, color, dim) in enumerate(topos):
    x = x0 + (card_w + gap) * i
    add_card(s, x, y, card_w, card_h)
    # status chip top
    add_round(s, x + Inches(0.3), y + Inches(0.3),
              Inches(1.4), Inches(0.32), dim, radius=0.5)
    add_oval(s, x + Inches(0.42), y + Inches(0.4),
             Inches(0.12), Inches(0.12), color)
    add_text(s, x + Inches(0.6), y + Inches(0.3),
             Inches(1.2), Inches(0.32), eyebrow,
             size=10, bold=True, color=color, font=FONT_MONO,
             anchor=MSO_ANCHOR.MIDDLE)
    # title
    add_text(s, x + Inches(0.3), y + Inches(0.85),
             card_w - Inches(0.6), Inches(0.5),
             title, size=18, bold=True, color=TEXT)
    # divider
    add_rect(s, x + Inches(0.3), y + Inches(1.45),
             card_w - Inches(0.6), Emu(6000), BORDER)
    # body
    add_text(s, x + Inches(0.3), y + Inches(1.6),
             card_w - Inches(0.6), card_h - Inches(1.8),
             body, size=12, color=TEXT_2, line_spacing=1.45)

# Bottom band — full-width promise
band_y = Inches(6.6)
add_card(s, Inches(0.55), band_y, Inches(12.2), Inches(0.45))
add_rect(s, Inches(0.55), band_y, Inches(0.04), Inches(0.45), TEAL)
add_text(s, Inches(0.85), band_y, Inches(12), Inches(0.45),
         "same install  ·  same workflow  ·  same UI  ·  same telemetry  "
         "·  for any repo onboarded with claudboard",
         size=11, color=TEXT_2, font=FONT_MONO,
         anchor=MSO_ANCHOR.MIDDLE)

footer(s, "12 / 13")


# =========================================================================
# Slide 13 — Recap / Q&A
# =========================================================================
s = blank_slide()

# Brand top-left
add_brand(s, Inches(0.6), Inches(0.6))
add_text(s, Inches(8.0), Inches(0.65), Inches(5.0), Inches(0.3),
         "END  ·  Q & A", size=10, bold=True, color=MUTED,
         font=FONT_MONO, align=PP_ALIGN.RIGHT)

# divider
add_rect(s, Inches(0.6), Inches(1.15), Inches(12.2), Emu(6000), BORDER)

# Title block
add_text(s, Inches(0.6), Inches(1.6), Inches(12), Inches(0.4),
         "RECAP", size=12, bold=True, color=TEAL, font=FONT_MONO)
add_text(s, Inches(0.6), Inches(2.0), Inches(12), Inches(1.2),
         "Two components.", size=48, bold=True, color=TEXT,
         line_spacing=1.0)
add_text(s, Inches(0.6), Inches(2.85), Inches(12), Inches(1.2),
         "One gate. Any repo.", size=48, bold=True, color=TEXT_2,
         line_spacing=1.0)

# Three pillars (mimic card grid)
y = Inches(4.1)
pcol_w = Inches(3.95)
pgap = Inches(0.2)
pillars = [
    ("01", "plugin", "Plugin",
     "claudboard onboards any repo:\nanalyse · generate · refresh · techdebt.",
     VIOLET, VIOLET_DIM),
    ("02", "workflow", "Workflow",
     "feature-workflow runs the lifecycle\nwith 10 sub-agents and 1 gate.",
     TEAL, TEAL_DIM),
    ("03", "surface", "Two surfaces",
     "Run from Claude Code CLI or\nfrom the Bosch SDLC dashboard.",
     AMBER, AMBER_DIM),
]
for i, (num, kind, h, b, color, dim) in enumerate(pillars):
    x = Inches(0.6) + (pcol_w + pgap) * i
    add_card(s, x, y, pcol_w, Inches(2.2))
    # chip
    add_round(s, x + Inches(0.3), y + Inches(0.3),
              Inches(0.9), Inches(0.3), dim, radius=0.5)
    add_text(s, x + Inches(0.3), y + Inches(0.3),
             Inches(0.9), Inches(0.3), kind,
             size=9, bold=True, color=color, font=FONT_MONO,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, x + pcol_w - Inches(0.7), y + Inches(0.3),
             Inches(0.5), Inches(0.3), num,
             size=11, bold=True, color=DIM, font=FONT_MONO,
             align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
    # divider
    add_rect(s, x + Inches(0.3), y + Inches(0.8),
             pcol_w - Inches(0.6), Emu(6000), BORDER)
    add_text(s, x + Inches(0.3), y + Inches(0.95),
             pcol_w - Inches(0.6), Inches(0.5), h,
             size=18, bold=True, color=TEXT)
    add_text(s, x + Inches(0.3), y + Inches(1.45),
             pcol_w - Inches(0.6), Inches(0.7), b,
             size=12, color=TEXT_2, line_spacing=1.4)

# Q&A row at bottom
qa_y = Inches(6.55)
add_text(s, Inches(0.6), qa_y, Inches(6), Inches(0.5),
         "Questions?", size=22, bold=True, color=TEXT)
# Try it pill
try_x = Inches(7.5); try_y = qa_y + Inches(0.07)
add_round(s, try_x, try_y, Inches(5.25), Inches(0.4),
          TEAL_DIM, radius=0.5)
add_text(s, try_x, try_y, Inches(5.25), Inches(0.4),
         "/analyse  →  /generate  →  /claudboard-workflow",
         size=11, bold=True, color=TEAL, font=FONT_MONO,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

# subtle footer line
add_rect(s, Inches(0.6), Inches(7.15), Inches(12.2), Emu(6000), BORDER)


prs.save(OUT)
print(f"Wrote {OUT}")
