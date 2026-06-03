#!/usr/bin/env python3
"""Add Measurement Framework slides to AI-SDLC-team-presentation-v2.pptx"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
import copy

# ── Design tokens ──────────────────────────────────────────
BG           = RGBColor(0x08, 0x09, 0x0A)
CARD_BG      = RGBColor(0x13, 0x14, 0x18)
RULE_COLOR   = RGBColor(0x23, 0x25, 0x2B)
TEAL         = RGBColor(0x4D, 0xD5, 0xDA)
PURPLE       = RGBColor(0xA7, 0x8B, 0xF0)
GOLD         = RGBColor(0xE3, 0xB0, 0x54)
GREEN        = RGBColor(0x4D, 0xDA, 0x80)
RED          = RGBColor(0xDA, 0x4D, 0x4D)
WHITE        = RGBColor(0xED, 0xED, 0xEF)
BODY         = RGBColor(0xB8, 0xB9, 0xBE)
MUTED        = RGBColor(0x79, 0x7B, 0x83)
DARK_MUTED   = RGBColor(0x4D, 0x4F, 0x57)
PILL_TEAL_BG = RGBColor(0x14, 0x28, 0x2A)
PILL_PURP_BG = RGBColor(0x20, 0x18, 0x2E)
PILL_GOLD_BG = RGBColor(0x2C, 0x21, 0x10)
PILL_GREEN_BG= RGBColor(0x14, 0x2A, 0x1A)
PILL_RED_BG  = RGBColor(0x2A, 0x14, 0x14)
LOGO_BG      = RGBColor(0xED, 0xED, 0xEF)

SLIDE_W = 12192000
SLIDE_H = 6858000

# Positions (from existing slides)
HEADER_Y     = Emu(384048)
RULE_Y       = Emu(868680)
SECTION_Y    = Emu(1097280)
TITLE_Y      = Emu(1371600)
SUBTITLE_Y   = Emu(1874519)
CONTENT_Y    = Emu(2560320)
FOOTER_RULE_Y= Emu(6446520)
FOOTER_Y     = Emu(6537960)
FOOTER_TEXT_Y= Emu(6519672)

LEFT_MARGIN  = Emu(502920)
RIGHT_EDGE   = Emu(11658600)
CONTENT_W    = Emu(11155680)

def add_bg(slide):
    rect = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Emu(12191695), Emu(SLIDE_H))
    rect.fill.solid()
    rect.fill.fore_color.rgb = BG
    rect.line.fill.background()

def add_header(slide):
    # "cb" badge
    badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, LEFT_MARGIN, HEADER_Y, Emu(292608), Emu(292608))
    badge.fill.solid()
    badge.fill.fore_color.rgb = LOGO_BG
    badge.line.fill.background()

    cb = slide.shapes.add_textbox(LEFT_MARGIN, HEADER_Y, Emu(292608), Emu(292608))
    tf = cb.text_frame
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "cb"
    r.font.name = "Menlo"
    r.font.size = Pt(11)
    r.font.bold = True
    r.font.color.rgb = BG
    tf.paragraphs[0].space_before = Pt(6)

    # "claudboard" text
    tb = slide.shapes.add_textbox(Emu(886968), Emu(402336), Emu(1828800), Emu(274320))
    p = tb.text_frame.paragraphs[0]
    r = p.add_run()
    r.text = "claudboard"
    r.font.name = "Inter"
    r.font.size = Pt(15)
    r.font.bold = True
    r.font.color.rgb = WHITE

    # Version badge
    ver_bg = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Emu(1920240), Emu(438912), Emu(502920), Emu(201168))
    ver_bg.fill.solid()
    ver_bg.fill.fore_color.rgb = BG
    ver_bg.line.color.rgb = RULE_COLOR

    ver = slide.shapes.add_textbox(Emu(1920240), Emu(429768), Emu(502920), Emu(201168))
    p = ver.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "V3.8.1"
    r.font.name = "Menlo"
    r.font.size = Pt(8)
    r.font.bold = True
    r.font.color.rgb = MUTED

    # Right-side path
    path = slide.shapes.add_textbox(Emu(7315200), Emu(438912), Emu(4754880), Emu(274320))
    p = path.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    r = p.add_run()
    r.text = "~/claude-repo-scan  ·  main"
    r.font.name = "Menlo"
    r.font.size = Pt(10)
    r.font.color.rgb = DARK_MUTED

    # Horizontal rule
    rule = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, LEFT_MARGIN, RULE_Y, CONTENT_W, Emu(6000))
    rule.fill.solid()
    rule.fill.fore_color.rgb = RULE_COLOR
    rule.line.fill.background()

def add_footer(slide, page_num, total):
    rule = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, LEFT_MARGIN, FOOTER_RULE_Y, CONTENT_W, Emu(6000))
    rule.fill.solid()
    rule.fill.fore_color.rgb = RULE_COLOR
    rule.line.fill.background()

    dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, LEFT_MARGIN, FOOTER_Y, Emu(91440), Emu(91440))
    dot.fill.solid()
    dot.fill.fore_color.rgb = TEAL
    dot.line.fill.background()

    ft = slide.shapes.add_textbox(Emu(658368), FOOTER_TEXT_Y, Emu(7315200), Emu(274320))
    p = ft.text_frame.paragraphs[0]
    r = p.add_run()
    r.text = "claude-repo-scan  ·  bosch-sdlc  ·  team review"
    r.font.name = "Inter"
    r.font.size = Pt(9)
    r.font.color.rgb = MUTED

    pn = slide.shapes.add_textbox(Emu(10058400), FOOTER_TEXT_Y, Emu(1828800), Emu(274320))
    p = pn.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    r = p.add_run()
    r.text = f"{page_num:02d} / {total:02d}"
    r.font.name = "Menlo"
    r.font.size = Pt(9)
    r.font.color.rgb = DARK_MUTED

def add_section_label(slide, text, color=TEAL):
    tb = slide.shapes.add_textbox(LEFT_MARGIN, SECTION_Y, Emu(7315200), Emu(274320))
    p = tb.text_frame.paragraphs[0]
    r = p.add_run()
    r.text = text
    r.font.name = "Inter"
    r.font.size = Pt(10)
    r.font.bold = True
    r.font.color.rgb = color

def add_title(slide, text):
    tb = slide.shapes.add_textbox(LEFT_MARGIN, TITLE_Y, CONTENT_W, Emu(640080))
    p = tb.text_frame.paragraphs[0]
    r = p.add_run()
    r.text = text
    r.font.name = "Inter"
    r.font.size = Pt(28)
    r.font.bold = True
    r.font.color.rgb = WHITE

def add_subtitle(slide, text):
    tb = slide.shapes.add_textbox(LEFT_MARGIN, SUBTITLE_Y, CONTENT_W, Emu(365760))
    p = tb.text_frame.paragraphs[0]
    r = p.add_run()
    r.text = text
    r.font.name = "Inter"
    r.font.size = Pt(13)
    r.font.color.rgb = MUTED

def add_card(slide, x, y, w, h, fill_color=CARD_BG):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    card.fill.solid()
    card.fill.fore_color.rgb = fill_color
    card.line.fill.background()
    return card

def add_text(slide, x, y, w, h, text, font_name="Inter", size=13, color=BODY, bold=False, align=PP_ALIGN.LEFT):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.name = font_name
    r.font.size = Pt(size)
    r.font.color.rgb = color
    r.font.bold = bold
    return tb

def add_bullet_list(slide, x, y, w, h, items, bullet_color=TEAL, text_color=BODY, size=13):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.space_after = Pt(4)
        r1 = p.add_run()
        r1.text = "—  "
        r1.font.name = "Menlo"
        r1.font.size = Pt(size)
        r1.font.color.rgb = bullet_color
        r2 = p.add_run()
        r2.text = item
        r2.font.name = "Inter"
        r2.font.size = Pt(size)
        r2.font.color.rgb = text_color
    return tb

def add_pill(slide, x, y, text, bg_color=PILL_TEAL_BG, text_color=TEAL, w=None):
    if w is None:
        w = Emu(len(text) * 91440 + 182880)
    pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, Emu(274320))
    pill.fill.solid()
    pill.fill.fore_color.rgb = bg_color
    pill.line.fill.background()

    tb = slide.shapes.add_textbox(x, y, w, Emu(274320))
    p = tb.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = text
    r.font.name = "Menlo"
    r.font.size = Pt(10)
    r.font.bold = True
    r.font.color.rgb = text_color

def add_card_rule(slide, x, y, w):
    rule = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, Emu(6000))
    rule.fill.solid()
    rule.fill.fore_color.rgb = RULE_COLOR
    rule.line.fill.background()

def add_numbered_card(slide, x, y, w, h, num, title, subtitle_text, items,
                      num_color=TEAL, num_bg=PILL_TEAL_BG, bullet_color=TEAL):
    """Card with number badge, title, rule, bullet list"""
    add_card(slide, x, y, w, h)

    inset_x = x + Emu(274320)
    inset_w = w - Emu(548640)

    # Number badge
    badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, inset_x, y + Emu(274320), Emu(365760), Emu(365760))
    badge.fill.solid()
    badge.fill.fore_color.rgb = num_bg
    badge.line.fill.background()

    add_text(slide, inset_x, y + Emu(274320), Emu(365760), Emu(365760),
             num, "Menlo", 12, num_color, True, PP_ALIGN.CENTER)

    # Title
    add_text(slide, inset_x + Emu(457200), y + Emu(292608), inset_w - Emu(457200), Emu(320040),
             title, "Inter", 18, WHITE, True)

    # Subtitle
    if subtitle_text:
        add_text(slide, inset_x + Emu(457200), y + Emu(548640), inset_w - Emu(457200), Emu(228600),
                 subtitle_text, "Menlo", 10, MUTED)

    # Rule
    rule_y = y + Emu(777240)
    add_card_rule(slide, inset_x, rule_y, inset_w)

    # Bullets
    if items:
        add_bullet_list(slide, inset_x, rule_y + Emu(137160), inset_w, h - Emu(914400),
                        items, bullet_color, BODY, 12)


# ── Load existing presentation ────────────────────────────
prs = Presentation("AI-SDLC-team-presentation-v2.pptx")
blank_layout = prs.slide_layouts[6]  # Blank layout

TOTAL_SLIDES = 23  # 13 existing + 10 new
start_num = 14

# ──────────────────────────────────────────────────────────
# SLIDE 14: Section intro — Measurement Framework
# ──────────────────────────────────────────────────────────
slide = prs.slides.add_slide(blank_layout)
add_bg(slide)
add_header(slide)

add_section_label(slide, "SECTION F  ·  MEASUREMENT FRAMEWORK")
add_title(slide, "Metrics for Agentic SDLC Adoption.")

add_subtitle(slide, "How to measure whether the human-agent delivery system is producing validated value — not just more AI activity.")

# Central question card
add_card(slide, LEFT_MARGIN, CONTENT_Y, CONTENT_W, Emu(1280160))
accent_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, LEFT_MARGIN, CONTENT_Y, Emu(36576), Emu(1280160))
accent_bar.fill.solid()
accent_bar.fill.fore_color.rgb = TEAL
accent_bar.line.fill.background()

add_text(slide, Emu(777240), Emu(2651760), Emu(10972800), Emu(274320),
         "CENTRAL QUESTION", "Inter", 10, TEAL, True)

add_text(slide, Emu(777240), Emu(2926080), Emu(10972800), Emu(457200),
         "Are we producing more validated value, with acceptable quality, cost, resilience, and learning?",
         "Inter", 20, WHITE, True)

# Six principles as 2x3 grid
principles = [
    ("01", "Measure validated value", "Not AI activity. Prompts submitted, LOC generated, and tickets closed are not success metrics."),
    ("02", "Separate output from outcome", "Producing more artifacts faster does not mean delivering more value."),
    ("03", "Segment by work type + risk", "A 70% autonomous rate is fine for config, dangerous for payments."),
    ("04", "Small operating scorecard", "Compact scorecard + diagnostic families. Not a reporting burden."),
    ("05", "Metrics for learning first", "Help teams learn where agents help, where humans are overloaded, where validation is weak."),
    ("06", "Balance five dimensions", "Speed only counts if quality, cost, resilience, and learning hold."),
]

for i, (num, title_text, desc) in enumerate(principles):
    col = i % 3
    row = i // 3
    cx = Emu(502920 + col * 3748920)
    cy = Emu(4114800 + row * 1188720)
    cw = Emu(3566040)
    ch = Emu(1005840)

    add_card(slide, cx, cy, cw, ch)

    add_text(slide, cx + Emu(182880), cy + Emu(137160), Emu(365760), Emu(274320),
             num, "Menlo", 11, TEAL, True)
    add_text(slide, cx + Emu(502920), cy + Emu(137160), cw - Emu(640080), Emu(274320),
             title_text, "Inter", 12, WHITE, True)
    add_text(slide, cx + Emu(182880), cy + Emu(457200), cw - Emu(365760), Emu(502920),
             desc, "Inter", 10, MUTED)

add_footer(slide, 14, TOTAL_SLIDES)


# ──────────────────────────────────────────────────────────
# SLIDE 15: Five Measurement Dimensions
# ──────────────────────────────────────────────────────────
slide = prs.slides.add_slide(blank_layout)
add_bg(slide)
add_header(slide)
add_section_label(slide, "CORE DEFINITIONS  ·  DIMENSIONS")
add_title(slide, "Five balancing dimensions.")
add_subtitle(slide, "No single metric should dominate. A healthy agentic SDLC must improve across all five.")

dims = [
    ("Speed", TEAL, PILL_TEAL_BG, "How quickly work moves from intent to validated change.", "Are we faster end to end, or only generating artifacts faster?"),
    ("Quality", PURPLE, PILL_PURP_BG, "Whether outputs are correct, maintainable, secure, and aligned with intent.", "Are agentic outputs reliable enough to scale?"),
    ("Cost", GOLD, PILL_GOLD_BG, "Total cost of producing validated change, including labor, tooling, review, rework.", "Are we reducing real cost or shifting cost elsewhere?"),
    ("Resilience", GREEN, PILL_GREEN_BG, "Ability to maintain control, recover from errors, avoid brittle dependence.", "Can the system absorb failures without losing control?"),
    ("Learning", RGBColor(0xDA, 0x8B, 0x4D), RGBColor(0x2A, 0x21, 0x14), "Whether people, teams, and workflows improve over time.", "Are we becoming more capable, or just more dependent?"),
]

for i, (name, color, bg_color, defn, question) in enumerate(dims):
    cx = LEFT_MARGIN
    cy = Emu(2560320 + i * 731520)
    cw = CONTENT_W
    ch = Emu(640080)

    add_card(slide, cx, cy, cw, ch)

    # Color accent bar on left
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, cx, cy, Emu(27432), ch)
    bar.fill.solid()
    bar.fill.fore_color.rgb = color
    bar.line.fill.background()

    # Dimension name
    add_text(slide, cx + Emu(274320), cy + Emu(91440), Emu(1600200), Emu(365760),
             name, "Inter", 16, color, True)

    # Definition
    add_text(slide, cx + Emu(1920240), cy + Emu(73152), Emu(5486400), Emu(502920),
             defn, "Inter", 12, BODY)

    # Question
    add_text(slide, cx + Emu(7772400), cy + Emu(73152), Emu(3200400), Emu(502920),
             question, "Inter", 11, MUTED)

add_footer(slide, 15, TOTAL_SLIDES)


# ──────────────────────────────────────────────────────────
# SLIDE 16: Validation Levels + Risk Classes
# ──────────────────────────────────────────────────────────
slide = prs.slides.add_slide(blank_layout)
add_bg(slide)
add_header(slide)
add_section_label(slide, "CORE DEFINITIONS  ·  VALIDATION & RISK")
add_title(slide, 'What "done" means — and how serious if wrong.')
add_subtitle(slide, "Validation levels prevent ambiguity. Risk classes prevent unsafe generalization of agent autonomy.")

# Left panel — Validation Levels
panel_w = Emu(5486400)
panel_h = Emu(3703320)
add_card(slide, LEFT_MARGIN, CONTENT_Y, panel_w, panel_h)
add_card_rule(slide, Emu(777240), Emu(3337560), Emu(4937760))

add_text(slide, Emu(777240), Emu(2743200), Emu(4937760), Emu(365760),
         "Validation Levels", "Inter", 20, WHITE, True)

v_levels = [
    ("V0", "Generated", "Artifact produced by human or agent."),
    ("V1", "Technically accepted", "Passed review, build, lint, tests."),
    ("V2", "Release-ready", "Passed all engineering, security, compliance checks."),
    ("V3", "Released", "Deployed or made available to users."),
    ("V4", "Outcome-validated", "Evidence of intended user/business outcome."),
]

for i, (code, name, desc) in enumerate(v_levels):
    vy = Emu(3474720 + i * 365760)
    add_text(slide, Emu(777240), vy, Emu(457200), Emu(274320),
             code, "Menlo", 11, TEAL, True)
    add_text(slide, Emu(1234440), vy, Emu(1828800), Emu(274320),
             name, "Inter", 12, WHITE, True)
    add_text(slide, Emu(3200400), vy, Emu(2514600), Emu(274320),
             desc, "Inter", 10, MUTED)

# Right panel — Risk Classes
add_card(slide, Emu(6172200), CONTENT_Y, panel_w, panel_h)
add_card_rule(slide, Emu(6446520), Emu(3337560), Emu(4937760))

add_text(slide, Emu(6446520), Emu(2743200), Emu(4937760), Emu(365760),
         "Risk Classes", "Inter", 20, WHITE, True)

r_classes = [
    ("R1", "Low", GREEN, "Limited blast radius, reversible.", "High autonomy acceptable"),
    ("R2", "Moderate", TEAL, "Some impact, testable and reversible.", "Agent + human review"),
    ("R3", "High", GOLD, "Significant customer/business/data impact.", "Tight human review"),
    ("R4", "Critical", RED, "Legal, compliance, financial, safety.", "Human-led or strictly gated"),
]

for i, (code, name, color, desc, autonomy) in enumerate(r_classes):
    ry = Emu(3474720 + i * 548640)
    add_text(slide, Emu(6446520), ry, Emu(457200), Emu(274320),
             code, "Menlo", 12, color, True)
    add_text(slide, Emu(6903720), ry, Emu(1280160), Emu(274320),
             name, "Inter", 13, WHITE, True)
    add_text(slide, Emu(6446520), ry + Emu(274320), Emu(4937760), Emu(228600),
             f"{desc}  →  {autonomy}", "Inter", 10, MUTED)

add_footer(slide, 16, TOTAL_SLIDES)


# ──────────────────────────────────────────────────────────
# SLIDE 17: Work Types
# ──────────────────────────────────────────────────────────
slide = prs.slides.add_slide(blank_layout)
add_bg(slide)
add_header(slide)
add_section_label(slide, "CORE DEFINITIONS  ·  WORK TYPES")
add_title(slide, "Six work types — segment everything.")
add_subtitle(slide, "Aggregated metrics are misleading. Work type describes what kind of engineering work is performed, not risk.")

wt_items = [
    ("WT1", "Configuration", TEAL, "Flags, CI/CD rules, access, build config.", "Fast to execute, can be high-risk"),
    ("WT2", "Customization", PURPLE, "Modify existing behavior, flow, rule, UI.", "Benefits from patterns, but business-rule risk can be significant"),
    ("WT3", "New Development", GOLD, "New capability, endpoint, component, flow.", "Higher ambiguity, needs strong intent + validation"),
    ("WT4", "Integration", GREEN, "Connect, sync, exchange with other systems.", "Risk in contracts, auth, data mapping, failure handling"),
    ("WT5", "Bug Fixing", RGBColor(0xDA, 0x8B, 0x4D), "Correct existing behavior that doesn't work.", "Agent suitability depends on root-cause clarity"),
    ("WT6", "Refactor / Tech Debt", RED, "Improve structure without behavior change.", "Requires strong regression tests and architectural context"),
]

for i, (code, name, color, desc, risk) in enumerate(wt_items):
    col = i % 3
    row = i // 3
    cx = Emu(502920 + col * 3748920)
    cy = Emu(2560320 + row * 1920240)
    cw = Emu(3566040)
    ch = Emu(1737360)

    add_card(slide, cx, cy, cw, ch)

    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, cx, cy, Emu(27432), ch)
    bar.fill.solid()
    bar.fill.fore_color.rgb = color
    bar.line.fill.background()

    add_text(slide, cx + Emu(274320), cy + Emu(182880), Emu(548640), Emu(320040),
             code, "Menlo", 12, color, True)
    add_text(slide, cx + Emu(822960), cy + Emu(182880), cw - Emu(1005840), Emu(320040),
             name, "Inter", 15, WHITE, True)

    add_card_rule(slide, cx + Emu(274320), cy + Emu(594360), cw - Emu(548640))

    add_text(slide, cx + Emu(274320), cy + Emu(731520), cw - Emu(548640), Emu(365760),
             desc, "Inter", 11, BODY)
    add_text(slide, cx + Emu(274320), cy + Emu(1188720), cw - Emu(548640), Emu(457200),
             risk, "Inter", 10, MUTED)

add_footer(slide, 17, TOTAL_SLIDES)


# ──────────────────────────────────────────────────────────
# SLIDE 18: Core Operating Scorecard
# ──────────────────────────────────────────────────────────
slide = prs.slides.add_slide(blank_layout)
add_bg(slide)
add_header(slide)
add_section_label(slide, "SCORECARD  ·  10 CORE METRICS")
add_title(slide, "The operating scorecard.")
add_subtitle(slide, "Organized around five dimensions. A small set that tells leadership where to look.")

scorecard = [
    ("Speed", "Lead time: intent → V2", "Flow", "Lagging", TEAL),
    ("Quality", "Change failure & rework burden", "Quality", "Lagging", PURPLE),
    ("Cost", "Cost per V2 validated change", "Economic", "Lagging", GOLD),
    ("Resilience", "Sustainable supervision load", "Risk", "Balancing", GREEN),
    ("Learning", "Prompt / intent quality improvement", "Capability", "Leading", RGBColor(0xDA, 0x8B, 0x4D)),
    ("Quality", "Defect escape rate", "Quality", "Lagging", PURPLE),
    ("Resilience", "Approval latency by risk class", "Flow/Risk", "Balancing", GREEN),
    ("Quality", "Agent rework rate by work type", "Quality", "Leading", PURPLE),
    ("Learning", "Technical understanding retention", "Capability", "Leading", RGBColor(0xDA, 0x8B, 0x4D)),
    ("Resilience", "Context freshness score", "Capability", "Leading", GREEN),
]

# Table header
header_y = Emu(2468880)
add_card(slide, LEFT_MARGIN, header_y, CONTENT_W, Emu(320040))
add_text(slide, Emu(777240), header_y + Emu(36576), Emu(1371600), Emu(256032),
         "DIMENSION", "Menlo", 9, MUTED, True)
add_text(slide, Emu(2468880), header_y + Emu(36576), Emu(4572000), Emu(256032),
         "CORE METRIC", "Menlo", 9, MUTED, True)
add_text(slide, Emu(7406640), header_y + Emu(36576), Emu(1554480), Emu(256032),
         "TYPE", "Menlo", 9, MUTED, True)
add_text(slide, Emu(9144000), header_y + Emu(36576), Emu(2468880), Emu(256032),
         "INDICATOR", "Menlo", 9, MUTED, True)

for i, (dim, metric, mtype, indicator, color) in enumerate(scorecard):
    ry = Emu(2834640 + i * 365760)

    if i % 2 == 0:
        stripe = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, LEFT_MARGIN, ry, CONTENT_W, Emu(365760))
        stripe.fill.solid()
        stripe.fill.fore_color.rgb = RGBColor(0x0E, 0x0F, 0x11)
        stripe.line.fill.background()

    # Color dot
    dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Emu(777240), ry + Emu(137160), Emu(91440), Emu(91440))
    dot.fill.solid()
    dot.fill.fore_color.rgb = color
    dot.line.fill.background()

    add_text(slide, Emu(960120), ry + Emu(54864), Emu(1371600), Emu(274320),
             dim, "Inter", 11, color, True)
    add_text(slide, Emu(2468880), ry + Emu(54864), Emu(4572000), Emu(274320),
             metric, "Inter", 12, WHITE)
    add_text(slide, Emu(7406640), ry + Emu(54864), Emu(1554480), Emu(274320),
             mtype, "Menlo", 10, MUTED)
    add_text(slide, Emu(9144000), ry + Emu(54864), Emu(2468880), Emu(274320),
             indicator, "Menlo", 10, MUTED)

add_footer(slide, 18, TOTAL_SLIDES)


# ──────────────────────────────────────────────────────────
# SLIDE 19: Diagnostic Metric Families
# ──────────────────────────────────────────────────────────
slide = prs.slides.add_slide(blank_layout)
add_bg(slide)
add_header(slide)
add_section_label(slide, "DIAGNOSTICS  ·  7 METRIC FAMILIES")
add_title(slide, "The scorecard says where — diagnostics say why.")
add_subtitle(slide, "Don't activate all at once. Use them when a core metric signals a problem.")

families = [
    ("Flow & Approval", TEAL, "When lead time or approval time worsens",
     "Cycle time, review queue, deployment readiness, flow efficiency, approval backlog"),
    ("Rework & Failure", PURPLE, "When quality or cost worsens",
     "Rework cost, defects after approval, regression frequency, rollback rate, root-cause clarity"),
    ("Agent Effectiveness", GOLD, "When deciding where to scale or restrict agents",
     "Accepted output ratio, autonomous completion, escalation rate, failure recurrence, traceability"),
    ("Supervision Load", GREEN, "When humans become the bottleneck",
     "Concurrent tasks/human, senior reviewer load, review fatigue, decision fatigue, reasoning reconstruction time"),
    ("Cost & Economics", RGBColor(0xDA, 0x8B, 0x4D), "When evaluating ROI or scaling",
     "AI cost per output, cost per feature, governance cost, delayed approval cost, workflow ROI"),
    ("Learning & Capability", RGBColor(0xDA, 0x4D, 0x8B), "When org wants to know if it's improving",
     "Workflow reuse, confidence calibration, validation maturity, context ownership, learning cycle closure"),
    ("Governance & Risk", RED, "When agentic workflows expand to regulated areas",
     "Bottleneck concentration, policy exceptions, security findings, unauthorized AI usage, audit trace completeness"),
]

for i, (name, color, trigger, metrics) in enumerate(families):
    cy = Emu(2468880 + i * 548640)

    add_card(slide, LEFT_MARGIN, cy, CONTENT_W, Emu(457200))

    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, LEFT_MARGIN, cy, Emu(27432), Emu(457200))
    bar.fill.solid()
    bar.fill.fore_color.rgb = color
    bar.line.fill.background()

    add_text(slide, Emu(777240), cy + Emu(45720), Emu(2286000), Emu(274320),
             name, "Inter", 13, color, True)
    add_text(slide, Emu(3200400), cy + Emu(45720), Emu(3200400), Emu(274320),
             trigger, "Inter", 11, WHITE)
    add_text(slide, Emu(3200400), cy + Emu(256032), Emu(8503920), Emu(228600),
             metrics, "Inter", 9, MUTED)

add_footer(slide, 19, TOTAL_SLIDES)


# ──────────────────────────────────────────────────────────
# SLIDE 20: Metrics by Maturity Stage
# ──────────────────────────────────────────────────────────
slide = prs.slides.add_slide(blank_layout)
add_bg(slide)
add_header(slide)
add_section_label(slide, "ADOPTION  ·  METRICS BY MATURITY STAGE")
add_title(slide, "Start small — expand with maturity.")
add_subtitle(slide, "Each stage has a recommended minimum viable metric set. Don't overload early adoption.")

stages = [
    ("AI-Assisted", "01", TEAL, PILL_TEAL_BG,
     "Humans drive work,\nAI assists with coding,\nexplanation, tests, docs.",
     ["Lead time → V2", "Defect escape rate", "Rework cost", "Test effectiveness",
      "Understanding retention", "Intent quality improvement"],
     "Don't prioritize yet: leverage ratio, autonomous completion, complex governance metrics."),
    ("HITL Agentic", "02", PURPLE, PILL_PURP_BG,
     "Agents produce artifacts,\nhumans still review closely.",
     ["Agent rework rate by WT", "Accepted output ratio", "Review queue time",
      "Human review effectiveness", "Approval latency by risk", "Cost per V2 change",
      "Context freshness score"],
     "Don't prioritize yet: org-wide ratio, full P3 calculations, complex confidence calibration."),
    ("Full Agentic", "03", GOLD, PILL_GOLD_BG,
     "Agents execute multi-step,\nhumans govern + approve\n+ handle exceptions.",
     ["Outcome-validated value", "Lead time → V2", "Supervision load",
      "Leverage ratio by risk", "Defects after approval", "Governance bottleneck",
      "Context freshness", "Understanding retention"],
     "Never use as headlines: % AI code, tasks launched, prompts submitted, raw autonomous rate."),
]

col_w = Emu(3566040)
for i, (name, num, color, bg_color, desc, metrics, caveat) in enumerate(stages):
    cx = Emu(502920 + i * 3748920)
    cy = CONTENT_Y
    ch = Emu(3749039)

    add_card(slide, cx, cy, col_w, ch)
    add_card_rule(slide, cx + Emu(274320), cy + Emu(777240), col_w - Emu(548640))

    # Pill
    add_pill(slide, cx + Emu(274320), cy + Emu(228600), name, bg_color, color, Emu(1371600))

    # Number
    add_text(slide, cx + col_w - Emu(640080), cy + Emu(228600), Emu(457200), Emu(274320),
             num, "Menlo", 11, DARK_MUTED, True, PP_ALIGN.RIGHT)

    # Description
    tb = slide.shapes.add_textbox(cx + Emu(274320), cy + Emu(914400), col_w - Emu(548640), Emu(640080))
    tf = tb.text_frame
    tf.word_wrap = True
    for line in desc.split("\n"):
        if tf.paragraphs[0].text == "":
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        r = p.add_run()
        r.text = line
        r.font.name = "Inter"
        r.font.size = Pt(11)
        r.font.color.rgb = MUTED

    # Metrics list
    my = cy + Emu(1645920)
    tb = slide.shapes.add_textbox(cx + Emu(274320), my, col_w - Emu(548640), Emu(1463040))
    tf = tb.text_frame
    tf.word_wrap = True
    for j, m in enumerate(metrics):
        if j == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.space_after = Pt(2)
        r1 = p.add_run()
        r1.text = "—  "
        r1.font.name = "Menlo"
        r1.font.size = Pt(10)
        r1.font.color.rgb = color
        r2 = p.add_run()
        r2.text = m
        r2.font.name = "Inter"
        r2.font.size = Pt(10)
        r2.font.color.rgb = BODY

    # Caveat
    add_text(slide, cx + Emu(274320), cy + ch - Emu(548640), col_w - Emu(548640), Emu(411480),
             caveat, "Inter", 9, DARK_MUTED)

add_footer(slide, 20, TOTAL_SLIDES)


# ──────────────────────────────────────────────────────────
# SLIDE 21: Interpretation Playbook
# ──────────────────────────────────────────────────────────
slide = prs.slides.add_slide(blank_layout)
add_bg(slide)
add_header(slide)
add_section_label(slide, "INTERPRETATION  ·  SIGNAL PATTERNS")
add_title(slide, "When metrics conflict — what it means.")
add_subtitle(slide, "Metrics should lead to decisions. These patterns help leaders and teams interpret conflicting signals.")

patterns = [
    ("Lead time ↓, defect escape ↑", "Speed bought with quality loss", "Strengthen validation, reduce autonomy for risky WT", RED),
    ("Agent rework ↑, intent quality ↓", "Agents getting unclear tasks", "Improve task framing, acceptance criteria, context", GOLD),
    ("Review queue ↑, agent cycle ↓", "Human review is the bottleneck", "Add capacity, automate low-risk, risk-based approvals", TEAL),
    ("Cost/V2 ↓, rework cost ↑", "Savings shifted downstream", "Include rework in cost model, inspect suitability", GOLD),
    ("Autonomous ↑, understanding ↓", "Human accountability weakening", "Add walkthroughs, explanation requirements, audits", RED),
    ("R1/R2 approval latency ↑", "Governance too heavy for low-risk", "Delegate or automate low-risk approvals", TEAL),
    ("R3/R4 latency ↓, defects ↑", "Governance too shallow", "Strengthen evidence requirements, expert review", RED),
    ("Context freshness ↓, failures ↑", "Agents using stale knowledge", "Assign context owners, update critical docs", GOLD),
]

# Column headers
hy = Emu(2377440)
add_text(slide, Emu(777240), hy, Emu(3200400), Emu(228600),
         "SIGNAL PATTERN", "Menlo", 9, MUTED, True)
add_text(slide, Emu(4114800), hy, Emu(3200400), Emu(228600),
         "LIKELY INTERPRETATION", "Menlo", 9, MUTED, True)
add_text(slide, Emu(7772400), hy, Emu(3886200), Emu(228600),
         "RECOMMENDED RESPONSE", "Menlo", 9, MUTED, True)

for i, (signal, interp, response, color) in enumerate(patterns):
    ry = Emu(2651760 + i * 457200)

    if i % 2 == 0:
        stripe = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, LEFT_MARGIN, ry, CONTENT_W, Emu(457200))
        stripe.fill.solid()
        stripe.fill.fore_color.rgb = RGBColor(0x0E, 0x0F, 0x11)
        stripe.line.fill.background()

    dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Emu(594360), ry + Emu(164592), Emu(91440), Emu(91440))
    dot.fill.solid()
    dot.fill.fore_color.rgb = color
    dot.line.fill.background()

    add_text(slide, Emu(777240), ry + Emu(91440), Emu(3200400), Emu(320040),
             signal, "Inter", 11, WHITE, True)
    add_text(slide, Emu(4114800), ry + Emu(91440), Emu(3200400), Emu(320040),
             interp, "Inter", 11, BODY)
    add_text(slide, Emu(7772400), ry + Emu(91440), Emu(3886200), Emu(320040),
             response, "Inter", 10, MUTED)

add_footer(slide, 21, TOTAL_SLIDES)


# ──────────────────────────────────────────────────────────
# SLIDE 22: Anti-Patterns
# ──────────────────────────────────────────────────────────
slide = prs.slides.add_slide(blank_layout)
add_bg(slide)
add_header(slide)
add_section_label(slide, "ANTI-PATTERNS  ·  WHAT NOT TO MEASURE", RED)
add_title(slide, "Metrics that will mislead you.")
add_subtitle(slide, "These metrics or behaviors should never be used as success indicators for agentic SDLC.")

anti_patterns = [
    ("% of AI-generated code", "Measures source of code, not value, quality, or maintainability."),
    ("Prompts per developer", "Encourages activity instead of outcomes."),
    ("Agent tasks launched", "Measures usage, not successful delivery."),
    ("Pull requests opened", "Agents can flood the review system with low-quality work."),
    ("Review speed alone", "Encourages shallow review and rubber-stamping."),
    ("Autonomous rate without risk segmentation", "Makes unsafe autonomy look like progress."),
    ("Individual AI productivity ranking", "Encourages gaming, fear, underreporting, poor collaboration."),
    ("Cost reduction without rework + governance", "Produces fake savings."),
]

for i, (metric, reason) in enumerate(anti_patterns):
    col = i % 2
    row = i // 2
    cx = Emu(502920 + col * 5669280)
    cy = Emu(2560320 + row * 822960)
    cw = Emu(5486400)
    ch = Emu(640080)

    add_card(slide, cx, cy, cw, ch)

    # Red X marker
    add_text(slide, cx + Emu(182880), cy + Emu(91440), Emu(365760), Emu(365760),
             "✗", "Menlo", 16, RED, True, PP_ALIGN.CENTER)

    add_text(slide, cx + Emu(548640), cy + Emu(91440), cw - Emu(731520), Emu(274320),
             metric, "Inter", 13, WHITE, True)
    add_text(slide, cx + Emu(548640), cy + Emu(365760), cw - Emu(731520), Emu(228600),
             reason, "Inter", 10, MUTED)

# Bottom callout
add_card(slide, LEFT_MARGIN, Emu(5852160), CONTENT_W, Emu(457200))
accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, LEFT_MARGIN, Emu(5852160), Emu(36576), Emu(457200))
accent.fill.solid()
accent.fill.fore_color.rgb = RED
accent.line.fill.background()

add_text(slide, Emu(777240), Emu(5897880), Emu(10972800), Emu(365760),
         "Never evaluate agentic SDLC using aggregate AI activity metrics. Use segmented, validated, risk-aware evidence.",
         "Inter", 13, WHITE, True)

add_footer(slide, 22, TOTAL_SLIDES)


# ──────────────────────────────────────────────────────────
# SLIDE 23: Pilot Decision Rules
# ──────────────────────────────────────────────────────────
slide = prs.slides.add_slide(blank_layout)
add_bg(slide)
add_header(slide)
add_section_label(slide, "DECISION  ·  SCALE / ADJUST / RESTRICT / STOP")
add_title(slide, "Pilot decision rules.")
add_subtitle(slide, "Every agentic workflow pilot should define baseline, comparison, and decision thresholds before rollout.")

decisions = [
    ("Scale", GREEN, PILL_GREEN_BG,
     "Lead time improves, quality stable or better, rework stable or lower, supervision sustainable, cost improves."),
    ("Adjust", GOLD, PILL_GOLD_BG,
     "Speed improves but rework, review load, or approval latency rises. Tune the workflow."),
    ("Restrict", RGBColor(0xDA, 0x8B, 0x4D), RGBColor(0x2A, 0x21, 0x14),
     "Quality, security, or understanding worsens in specific work types or risk classes. Limit to lower-risk work."),
    ("Stop", RED, PILL_RED_BG,
     "Total cost rises and quality worsens after multiple measurement cycles. Full stop and reassess."),
]

for i, (name, color, bg_color, desc) in enumerate(decisions):
    cx = Emu(502920 + i * 2834640)
    cy = CONTENT_Y
    cw = Emu(2651760)
    ch = Emu(2560320)

    add_card(slide, cx, cy, cw, ch)

    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, cx, cy, Emu(27432), ch)
    bar.fill.solid()
    bar.fill.fore_color.rgb = color
    bar.line.fill.background()

    # Arrow icon
    icon_bg = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx + Emu(274320), cy + Emu(274320), Emu(457200), Emu(457200))
    icon_bg.fill.solid()
    icon_bg.fill.fore_color.rgb = bg_color
    icon_bg.line.fill.background()

    icons = {"Scale": "↑", "Adjust": "↻", "Restrict": "▼", "Stop": "✖"}
    add_text(slide, cx + Emu(274320), cy + Emu(274320), Emu(457200), Emu(457200),
             icons[name], "Menlo", 18, color, True, PP_ALIGN.CENTER)

    add_text(slide, cx + Emu(822960), cy + Emu(365760), cw - Emu(1005840), Emu(365760),
             name, "Inter", 20, WHITE, True)

    add_card_rule(slide, cx + Emu(274320), cy + Emu(868680), cw - Emu(548640))

    add_text(slide, cx + Emu(274320), cy + Emu(1005840), cw - Emu(548640), Emu(1371600),
             desc, "Inter", 12, BODY)

# Bottom panel — Baseline requirements
add_card(slide, LEFT_MARGIN, Emu(5349240), CONTENT_W, Emu(960120))
accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, LEFT_MARGIN, Emu(5349240), Emu(36576), Emu(960120))
accent.fill.solid()
accent.fill.fore_color.rgb = TEAL
accent.line.fill.background()

add_text(slide, Emu(777240), Emu(5394960), Emu(10972800), Emu(274320),
         "BASELINE BEFORE SCALING", "Inter", 10, TEAL, True)

add_text(slide, Emu(777240), Emu(5669280), Emu(10972800), Emu(548640),
         "Establish current lead time, defect/rework levels, review load, cost per change, quality expectations, and validation level. "
         "Never compare simple agentic work with complex historical human work, or V1 acceptance with V4 outcome validation.",
         "Inter", 11, BODY)

add_footer(slide, 23, TOTAL_SLIDES)

# ── Update existing slide page numbers to reflect new total ──
# Existing slides had "NN / 13" — update to "NN / 23"
for i, slide in enumerate(prs.slides):
    if i >= 13:
        break
    for shape in slide.shapes:
        if shape.has_text_frame:
            for para in shape.text_frame.paragraphs:
                full_text = ''.join(r.text for r in para.runs)
                if f"/ 13" in full_text:
                    for run in para.runs:
                        if "/ 13" in run.text:
                            run.text = run.text.replace("/ 13", f"/ {TOTAL_SLIDES}")

# ── Save ──────────────────────────────────────────────────
output = "AI-SDLC-team-presentation-v2.pptx"
prs.save(output)
print(f"Saved {output} with {len(prs.slides)} slides")
